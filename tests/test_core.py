"""离线核心测试：不访问用户账号、浏览器会话或哔哩哔哩网站。"""
import asyncio
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

import db
from bilibili_client import BilibiliBrowserClient, parse_bvid, parse_mid
import downloader
from downloader import author_download_folder


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.original_path = db.DB_PATH
        db.DB_PATH = Path(self.folder.name) / "bili-test.db"
        db.init_db()

    def tearDown(self):
        db.DB_PATH = self.original_path
        self.folder.cleanup()

    def test_parsers_and_view_mapping(self):
        self.assertEqual(parse_bvid("分享 https://www.bilibili.com/video/BV1xx411c7mD"), "BV1xx411c7mD")
        self.assertEqual(parse_mid("https://space.bilibili.com/123456"), "123456")
        video = BilibiliBrowserClient._video_from_view({
            "bvid": "BV1xx411c7mD", "aid": 7, "title": "<em>测试</em> 标题", "pubdate": 1700000000,
            "owner": {"name": "测试UP", "mid": 8},
            "stat": {"view": 1, "like": 2, "reply": 3, "share": 4, "favorite": 5},
            "pages": [{"cid": 9, "page": 1, "part": "第一P"}],
        })
        self.assertEqual(video["title"], "测试 标题")
        self.assertEqual((video["view_count"], video["like_count"], video["comment_count"], video["share_count"], video["favorite_count"]), (1, 2, 3, 4, 5))
        self.assertEqual(video["pages"][0]["part"], "第一P")

    def test_result_storage_and_xlsx(self):
        db.upsert_result("single", {
            "bvid": "BV1xx411c7mD", "aid": "7", "author_name": "测试UP", "author_mid": "8",
            "title": "测试标题", "video_url": "https://www.bilibili.com/video/BV1xx411c7mD", "published_at": 1700000000,
            "view_count": 10, "like_count": 20, "comment_count": 30, "share_count": 40, "favorite_count": 50,
            "pages": [{"cid": 1, "page": 1, "part": "P1"}],
        }, input_url="https://www.bilibili.com/video/BV1xx411c7mD")
        rows = db.list_results("single")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["pages"][0]["cid"], 1)
        from web import app as web_app
        payload = asyncio.run(web_app.api_single_results())
        self.assertEqual(payload["rows"][0]["bvid"], "BV1xx411c7mD")
        workbook = load_workbook(BytesIO(web_app._xlsx_bytes("single")))
        worksheet = workbook.active
        self.assertIn("作者ID", [cell.value for cell in worksheet[1]])
        self.assertIn("评论数", [cell.value for cell in worksheet[1]])
        self.assertEqual(worksheet.cell(2, 2).value, "8")

    def test_single_result_keeps_duplicate_bvids_in_input_order(self):
        first = {
            "bvid": "BV1xx411c7mD", "aid": "7", "author_name": "测试UP", "author_mid": "8",
            "title": "第一次", "video_url": "https://www.bilibili.com/video/BV1xx411c7mD", "published_at": 1700000000,
            "view_count": 1, "like_count": 2, "comment_count": 3, "share_count": 4, "favorite_count": 5, "pages": [],
        }
        second = dict(first, title="第二次")
        db.upsert_result("single", first, input_url="first", input_order=0)
        db.upsert_result("single", second, input_url="second", input_order=1)
        rows = db.list_results("single")
        self.assertEqual([row["input_url"] for row in rows], ["first", "second"])
        self.assertEqual([row["title"] for row in rows], ["第一次", "第二次"])

    def test_creator_results_and_xlsx_follow_added_author_order(self):
        from web import app as web_app

        db.add_creator("1001", "https://space.bilibili.com/1001")
        db.add_creator("1002", "https://space.bilibili.com/1002")
        db.update_creator_profile("1001", {"nickname": "Z作者"})
        db.update_creator_profile("1002", {"nickname": "A作者"})
        for bvid, mid, name, published in [
            ("BV1A", "1002", "A作者", 30),
            ("BV1Zold", "1001", "Z作者", 10),
            ("BV1Znew", "1001", "Z作者", 20),
        ]:
            db.upsert_result("creator", {
                "bvid": bvid, "author_mid": mid, "author_name": name,
                "title": bvid, "published_at": published, "pages": [],
            })

        self.assertEqual(
            [row["bvid"] for row in db.list_results("creator")],
            ["BV1Znew", "BV1Zold", "BV1A"],
        )
        worksheet = load_workbook(BytesIO(web_app._xlsx_bytes("creator"))).active
        self.assertEqual(
            [worksheet.cell(row, 1).value for row in range(2, 5)],
            ["Z作者", "Z作者", "A作者"],
        )

    def test_single_link_normalization_keeps_duplicate_lines(self):
        from web import app as web_app
        links, invalid = web_app._normalise_lines(
            "https://www.bilibili.com/video/BV1xx411c7mD\nhttps://b23.tv/abcd\nhttps://www.bilibili.com/video/BV1xx411c7mD",
            web_app.SINGLE_URL_PATTERN,
            deduplicate=False,
        )
        self.assertEqual(invalid, [])
        self.assertEqual(links, [
            "https://www.bilibili.com/video/BV1xx411c7mD",
            "https://b23.tv/abcd",
            "https://www.bilibili.com/video/BV1xx411c7mD",
        ])

    def test_single_video_enters_bilibili_before_fetching_direct_bvid(self):
        class FakePage:
            def __init__(self):
                self.url = "about:blank"
                self.visits = []

            async def goto(self, url, **_kwargs):
                self.url = url
                self.visits.append(url)

        class FakeContext:
            async def close(self):
                pass

        class FakePlaywright:
            async def stop(self):
                pass

        page = FakePage()
        client = BilibiliBrowserClient()

        async def fake_launch():
            return FakePlaywright(), FakeContext(), page

        async def fake_view(current_page, bvid):
            self.assertEqual(current_page.url, "https://www.bilibili.com/")
            return {"bvid": bvid}

        client._launch = fake_launch
        client._view = fake_view
        videos, failures = asyncio.run(client.fetch_single_videos([
            "https://www.bilibili.com/video/BV1eUty6xErd"
        ]))
        self.assertEqual(failures, [])
        self.assertEqual(videos, [{"bvid": "BV1eUty6xErd"}])
        self.assertEqual(page.visits, ["https://www.bilibili.com/"])

    def test_single_invalid_link_is_rejected_before_browser(self):
        from web import app as web_app
        response = asyncio.run(web_app.api_single_extract({"urls": "not a bilibili link"}))
        self.assertEqual(response.status_code, 400)

    def test_author_download_folder_uses_safe_name_and_unique_mid(self):
        self.assertEqual(
            author_download_folder({"author_name": "同名/UP", "author_mid": "12345", "bvid": "BV1abc"}),
            "同名_UP（12345）",
        )
        self.assertEqual(
            author_download_folder({"author_name": "", "author_mid": "12345", "bvid": "BV1abc"}),
            "12345",
        )
        self.assertEqual(
            author_download_folder({"author_name": "", "author_mid": "", "bvid": "BV1abc"}),
            "未知作者（BV1abc）",
        )

    def test_download_template_numbers_single_and_multi_part_videos(self):
        """Replacing the fallback with raw playlist_index would regress single-part names to PNA."""
        from yt_dlp import YoutubeDL

        template = getattr(
            downloader,
            "download_output_template",
            lambda _title: "%(title)s - P%(playlist_index)02d.%(ext)s",
        )("测试作品")
        self.assertEqual(
            YoutubeDL({"outtmpl": template}).prepare_filename({"title": "测试作品", "ext": "mp4"}),
            "测试作品 - P01.mp4",
        )
        self.assertEqual(
            YoutubeDL({"outtmpl": template}).prepare_filename({"title": "测试作品", "ext": "mp4", "playlist_index": 3}),
            "测试作品 - P03.mp4",
        )


if __name__ == "__main__":
    unittest.main()
