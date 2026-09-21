"""哔哩哔哩视频提取 - 本地 FastAPI 界面。"""
import asyncio
import concurrent.futures
import logging
import re
import shutil
import threading
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path

from fastapi import Body, FastAPI, Request
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from jinja2 import Environment, FileSystemLoader

import db
from bilibili_client import BilibiliBrowserClient, parse_mid
from config import SESSION_DIR
from downloader import start_download

logger = logging.getLogger(__name__)
BASE_DIR = Path(__file__).parent
env = Environment(loader=FileSystemLoader(str(BASE_DIR / "templates")))


@asynccontextmanager
async def lifespan(app: FastAPI):
    db.init_db()
    yield


app = FastAPI(title="哔哩哔哩视频提取", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
_login_active = False
_creator_state = {"running": False, "current": 0, "total": 0, "found": 0, "name": "", "message": ""}
_topic_state = {"running": False, "current": 0, "total": 0, "found": 0, "label": "", "message": ""}
_single_state = {"running": False, "current": 0, "total": 0, "found": 0, "failed": 0, "message": ""}
_download_state = {"running": False, "total": 0, "done": 0, "failed": 0, "message": ""}
SINGLE_URL_PATTERN = re.compile(
    r"https?://(?:(?:www\.)?bilibili\.com/[^\s]+|b23\.tv/[^\s]+)", re.IGNORECASE
)


def render(name: str) -> HTMLResponse:
    return HTMLResponse(env.get_template(name).render())


def _fmt_time(value: int | None) -> str:
    if not value:
        return "—"
    return datetime.fromtimestamp(int(value)).strftime("%Y-%m-%d %H:%M")


def _parse_range(start_at: str, end_at: str, *, required: bool) -> tuple[int | None, int | None]:
    if required and (not start_at or not end_at):
        raise ValueError("请选择完整的开始和结束时间")
    if not start_at and not end_at:
        return None, None
    try:
        start = int(datetime.fromisoformat(start_at).timestamp()) if start_at else None
        end = int(datetime.fromisoformat(end_at).timestamp()) if end_at else None
    except ValueError as exc:
        raise ValueError("日期格式无效，请精确到分钟") from exc
    if start is not None and end is not None and start > end:
        raise ValueError("开始时间不能晚于结束时间")
    return start, end


def _is_busy() -> bool:
    return bool(_creator_state["running"] or _topic_state["running"] or _single_state["running"])


def _normalise_lines(value: str, pattern: re.Pattern, *, deduplicate: bool = True) -> tuple[list[str], list[str]]:
    values, invalid = [], []
    for line in (line.strip() for line in value.splitlines() if line.strip()):
        match = pattern.search(line)
        if match:
            values.append(match.group(0).rstrip("，。；：、）】》'\""))
        else:
            invalid.append(line)
    return list(dict.fromkeys(values)) if deduplicate else values, invalid


def _result_rows(source: str) -> list[dict]:
    rows = db.list_results(source)
    for row in rows:
        row["publish_time"] = _fmt_time(row["published_at"])
    return rows


@app.get("/", response_class=HTMLResponse)
async def creator_page(request: Request):
    return render("creator.html")


@app.get("/topic", response_class=HTMLResponse)
async def topic_page(request: Request):
    return render("topic.html")


@app.get("/single", response_class=HTMLResponse)
async def single_page(request: Request):
    return render("single.html")


@app.post("/api/login")
async def api_login():
    global _login_active
    if _login_active:
        return {"message": "登录窗口已打开，请完成扫码"}
    _login_active = True

    def _login():
        global _login_active
        try:
            asyncio.run(BilibiliBrowserClient().open_login())
        except Exception:
            logger.exception("打开哔哩哔哩登录窗口失败")
        finally:
            _login_active = False

    threading.Thread(target=_login, daemon=True).start()
    return {"message": "已打开哔哩哔哩登录窗口，请扫码登录"}


@app.delete("/api/browser-cache")
async def api_clear_browser_cache():
    if _login_active or _is_busy():
        return JSONResponse({"error": "请先完成当前登录或提取任务"}, 409)
    try:
        if SESSION_DIR.exists():
            shutil.rmtree(SESSION_DIR)
        SESSION_DIR.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        return JSONResponse({"error": f"清空浏览器缓存失败：{exc}"}, 500)
    return {"message": "已清空本程序的哔哩哔哩登录态和浏览器缓存"}


@app.get("/api/creators")
async def api_creators():
    return {"creators": db.list_creators()}


@app.delete("/api/creators")
async def api_clear_creators():
    if _is_busy():
        return JSONResponse({"error": "提取进行中，暂不能清空"}, 409)
    return db.clear_creators()


@app.post("/api/creators/add-profile")
async def api_add_profiles(payload: dict = Body(...)):
    if _is_busy():
        return JSONResponse({"error": "提取进行中，暂不能读取资料"}, 409)
    lines = [line.strip() for line in str(payload.get("urls", "")).splitlines() if line.strip()]
    if not lines:
        return JSONResponse({"error": "请逐行粘贴至少一个 UP 主空间链接"}, 400)
    mids, failed = [], []
    for line in lines:
        mid = parse_mid(line)
        if not mid:
            failed.append({"input": line, "error": "无法识别 UP 主空间链接"})
            continue
        db.add_creator(mid, line)
        mids.append(mid)
    if mids:
        def _read():
            return asyncio.run(BilibiliBrowserClient().read_profiles(mids))
        profiles = await asyncio.get_running_loop().run_in_executor(None, _read)
        for mid in mids:
            profile, error = profiles.get(mid, (None, "未读取到资料"))
            if profile:
                db.update_creator_profile(mid, profile)
            else:
                db.update_creator_profile(mid, {}, status="资料未读取")
                failed.append({"input": mid, "error": error})
    return {"added": [item for item in db.list_creators() if item["mid"] in mids], "failed": failed}


@app.post("/api/creators/refresh-unread")
async def api_refresh_unread():
    unread = [item for item in db.list_creators() if item.get("profile_status") != "已读取"]
    if not unread:
        return {"total": 0, "refreshed": 0, "failed": []}
    def _read():
        return asyncio.run(BilibiliBrowserClient().read_profiles([item["mid"] for item in unread]))
    profiles = await asyncio.get_running_loop().run_in_executor(None, _read)
    refreshed, failed = 0, []
    for item in unread:
        profile, error = profiles.get(item["mid"], (None, "未读取到资料"))
        if profile:
            db.update_creator_profile(item["mid"], profile)
            refreshed += 1
        else:
            failed.append({"input": item["mid"], "error": error})
    return {"total": len(unread), "refreshed": refreshed, "failed": failed}


@app.get("/api/creator-status")
async def api_creator_status():
    return _creator_state


@app.get("/api/creator-results")
async def api_creator_results():
    return {"rows": _result_rows("creator")}


@app.post("/api/creator-extract")
async def api_creator_extract(payload: dict = Body(...)):
    if _is_busy():
        return JSONResponse({"error": "已有提取任务正在运行，请等待完成"}, 409)
    creators = db.list_creators()
    if not creators:
        return JSONResponse({"error": "请先添加 UP 主空间链接"}, 400)
    top_count = payload.get("top_count")
    try:
        top_count = int(top_count) if top_count not in (None, "") else None
    except (ValueError, TypeError):
        return JSONResponse({"error": "前几条必须为正整数"}, 400)
    if top_count is not None and not 1 <= top_count <= 1000:
        return JSONResponse({"error": "前几条需在 1 到 1000 之间"}, 400)
    start_at, end_at = str(payload.get("start_at", "")), str(payload.get("end_at", ""))
    if top_count is not None:
        start_at = end_at = ""
    try:
        start_ts, end_ts = _parse_range(start_at, end_at, required=top_count is None)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    cleared = db.clear_results("creator")
    _creator_state.update({"running": True, "current": 0, "total": len(creators), "found": 0, "name": "", "message": "准备打开 UP 主空间"})

    def _crawl_all():
        failures = []
        for index, creator in enumerate(creators, start=1):
            _creator_state.update({"current": index, "name": creator.get("nickname") or creator["mid"], "message": "正在读取投稿视频"})
            def _save(video):
                db.upsert_result("creator", video, input_url=creator.get("input_url", ""))
                _creator_state["found"] += 1
                _creator_state["message"] = f"已即时显示 {_creator_state['found']} 条作品"
            try:
                asyncio.run(BilibiliBrowserClient().fetch_creator_videos(
                    creator["mid"], start_ts=start_ts, end_ts=end_ts, top_count=top_count, on_video=_save,
                ))
            except Exception as exc:
                failures.append({"mid": creator["mid"], "error": str(exc)})
        return failures

    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
            failures = await asyncio.get_running_loop().run_in_executor(pool, _crawl_all)
        _creator_state["message"] = f"提取完成：{_creator_state['found']} 条作品"
        return {"total": _creator_state["found"], "failed": failures, "cleared": cleared}
    except Exception as exc:
        logger.exception("UP主作品提取失败")
        _creator_state["message"] = f"提取失败：{exc}"
        return JSONResponse({"error": _creator_state["message"]}, 502)
    finally:
        _creator_state["running"] = False


@app.get("/api/topic-status")
async def api_topic_status():
    return _topic_state


@app.get("/api/topic-results")
async def api_topic_results():
    return {"rows": _result_rows("topic")}


@app.post("/api/topic-extract")
async def api_topic_extract(payload: dict = Body(...)):
    if _is_busy():
        return JSONResponse({"error": "已有提取任务正在运行，请等待完成"}, 409)
    raw_topic = str(payload.get("topic", "")).strip()
    if not raw_topic:
        return JSONResponse({"error": "请输入关键词或哔哩哔哩话题链接"}, 400)
    try:
        max_results = int(payload.get("max_results"))
    except (ValueError, TypeError):
        return JSONResponse({"error": "最多提取条数必须为正整数"}, 400)
    if not 1 <= max_results <= 1000:
        return JSONResponse({"error": "最多提取条数需在 1 到 1000 之间"}, 400)
    try:
        start_ts, end_ts = _parse_range(str(payload.get("start_at", "")), str(payload.get("end_at", "")), required=True)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    cleared = db.clear_results("topic")
    _topic_state.update({"running": True, "current": 0, "total": max_results, "found": 0, "label": raw_topic, "message": "准备打开搜索或话题页面"})

    def _crawl():
        label = raw_topic
        def _save(video):
            db.upsert_result("topic", video, query_label=label, input_url=raw_topic)
            _topic_state["found"] += 1
            _topic_state["current"] = _topic_state["found"]
            _topic_state["message"] = f"已即时显示 {_topic_state['found']} 条符合条件的视频"
        label, videos = asyncio.run(BilibiliBrowserClient().fetch_search_videos(
            raw_topic, start_ts=start_ts, end_ts=end_ts, max_results=max_results, on_video=_save,
        ))
        _topic_state["label"] = label
        return videos
    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
            videos = await asyncio.get_running_loop().run_in_executor(pool, _crawl)
        _topic_state["message"] = f"提取完成：{len(videos)} 条符合条件的视频"
        return {"total": len(videos), "cleared": cleared}
    except Exception as exc:
        _topic_state["message"] = f"提取失败：{exc}"
        return JSONResponse({"error": _topic_state["message"]}, 502)
    finally:
        _topic_state["running"] = False


@app.get("/api/single-status")
async def api_single_status():
    return _single_state


@app.get("/api/single-results")
async def api_single_results():
    return {"rows": _result_rows("single")}


@app.post("/api/single-extract")
async def api_single_extract(payload: dict = Body(...)):
    if _is_busy():
        return JSONResponse({"error": "已有提取任务正在运行，请等待完成"}, 409)
    urls, invalid = _normalise_lines(str(payload.get("urls", "")), SINGLE_URL_PATTERN, deduplicate=False)
    if not urls or invalid:
        return JSONResponse({"error": "请逐行粘贴有效的哔哩哔哩视频链接"}, 400)
    if len(urls) > 1000:
        return JSONResponse({"error": "一次最多提取 1000 个作品链接"}, 400)
    cleared = db.clear_results("single")
    _single_state.update({"running": True, "current": 0, "total": len(urls), "found": 0, "failed": 0, "message": "准备打开第一个视频链接"})
    def _crawl():
        def _save(position, input_url, video):
            db.upsert_result("single", video, input_url=input_url, input_order=position - 1)
            _single_state["found"] += 1
            _single_state["message"] = f"已即时显示 {_single_state['found']} 条视频"
        def _progress(current, _, error):
            _single_state["current"] = current
            if error:
                _single_state["failed"] += 1
                _single_state["message"] = f"第 {current} 条未读取到资料，正在继续"
        return asyncio.run(BilibiliBrowserClient().fetch_single_videos(urls, on_video=_save, on_progress=_progress))
    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
            videos, failures = await asyncio.get_running_loop().run_in_executor(pool, _crawl)
        _single_state["message"] = f"完成：成功 {len(videos)} 条，失败 {len(failures)} 条"
        return {"total": len(videos), "failed": failures, "cleared": cleared}
    except Exception as exc:
        logger.exception("单作品提取失败")
        _single_state["message"] = f"提取失败：{exc}"
        return JSONResponse({"error": _single_state["message"]}, 502)
    finally:
        _single_state["running"] = False


def _xlsx_bytes(source: str) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    labels = {"creator": "UP主提取结果", "topic": "关键词话题提取结果", "single": "单作品提取结果"}
    rows = _result_rows(source)
    wb = Workbook()
    ws = wb.active
    ws.title = labels[source]
    headers = ["UP主", "作者ID", "发布时间", "标题", "播放量", "点赞数", "评论数", "分享数", "收藏数", "作品链接"]
    ws.append(headers)
    for row in rows:
        ws.append([row["author_name"], row["author_mid"], row["publish_time"], row["title"], row["view_count"], row["like_count"],
                   row["comment_count"], row["share_count"], row["favorite_count"], row["video_url"]])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00A1D6")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in {"A": 20, "B": 16, "C": 20, "D": 50, "E": 14, "F": 14, "G": 14, "H": 14, "I": 14, "J": 48}.items():
        ws.column_dimensions[column].width = width
    for row_index in range(2, ws.max_row + 1):
        for column_index in range(5, 10):
            ws.cell(row_index, column_index).number_format = "#,##0"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


@app.post("/api/results/{source}/export-xlsx/save")
async def api_export_xlsx(source: str):
    if source not in {"creator", "topic", "single"}:
        return JSONResponse({"error": "无效结果类型"}, 404)
    if not db.list_results(source):
        return JSONResponse({"error": "暂无可导出的结果"}, 400)
    try:
        import tkinter as tk
        from tkinter import filedialog
        dialog = tk.Tk(); dialog.withdraw(); dialog.attributes("-topmost", True)
        names = {"creator": "哔哩哔哩UP主提取.xlsx", "topic": "哔哩哔哩话题提取.xlsx", "single": "哔哩哔哩单作品提取.xlsx"}
        output = filedialog.asksaveasfilename(parent=dialog, title="保存提取结果", defaultextension=".xlsx",
                                               initialfile=names[source], filetypes=[("Excel 工作簿", "*.xlsx")])
        dialog.destroy()
        if not output:
            return {"cancelled": True}
        Path(output).write_bytes(_xlsx_bytes(source))
        return {"ok": True, "path": output}
    except Exception as exc:
        return JSONResponse({"error": f"另存为失败：{exc}"}, 500)


@app.get("/api/download-status")
async def api_download_status():
    return _download_state


@app.post("/api/downloads/start")
async def api_download(payload: dict = Body(...)):
    if _download_state["running"]:
        return JSONResponse({"error": "已有下载任务正在进行"}, 409)
    raw_ids = payload.get("ids") or []
    if not isinstance(raw_ids, list):
        return JSONResponse({"error": "下载参数无效"}, 400)
    try:
        ids = list(dict.fromkeys(int(value) for value in raw_ids))
    except (TypeError, ValueError):
        return JSONResponse({"error": "作品标识无效"}, 400)
    items = db.get_results_by_ids(ids)
    if not items:
        return JSONResponse({"error": "请先勾选要下载的视频"}, 400)
    try:
        import tkinter as tk
        from tkinter import filedialog
        dialog = tk.Tk(); dialog.withdraw(); dialog.attributes("-topmost", True)
        output = filedialog.askdirectory(parent=dialog, title="选择视频保存文件夹", mustexist=True)
        dialog.destroy()
        if not output:
            return {"cancelled": True}
        start_download(items, Path(output), _download_state)
        return {"ok": True, "total": len(items), "path": output}
    except Exception as exc:
        return JSONResponse({"error": f"选择保存位置失败：{exc}"}, 500)
